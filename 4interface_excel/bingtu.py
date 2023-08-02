# /usr/bin/env python
# -*- coding: utf-8 -*-
# author__ = 'Aaron Han'
from openpyxl import load_workbook
from openpyxl.chart import PieChart, Reference


def pie_chart(wp):
    table = wp.get_sheet_by_name(wp.get_sheet_names()[0])

    # 生成饼图对象
    pie = PieChart()
    # 图的标题
    pie.title = "API接口测试统计"
    '''行数和列数都是从1开始的，和遍历用例是一样都是从1开始'''
    # 获取标签(取范围第一列的最小行数和最大行数)
    labels = Reference(table, min_col=4, min_row=6, max_col=4, max_row=7)
    # 获取数据（取范围第二列的最小行数-1和最大行数）
    data = Reference(table, min_col=5, min_row=5, max_col=5, max_row=7)

    # 添加数据和标签到饼图中
    pie.add_data(data, titles_from_data=True)
    pie.set_categories(labels)

    # 放在excel表中
    table.add_chart(pie, "A10")

    # 保存excel
    # wb.save("test1.xlsx")

# wb = load_workbook("D:\\PycharmProjects\\test.xlsx")
# pie_chart(wb)
